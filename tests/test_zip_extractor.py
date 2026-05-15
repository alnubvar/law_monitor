from __future__ import annotations

import io
import unittest
import zipfile

from app.extractors.zip_extractor import (
    MAX_ZIP_DOWNLOAD_BYTES,
    MAX_ZIP_UNCOMPRESSED_BYTES,
    extract_text_from_zip,
)
from unittest.mock import patch, MagicMock


def _make_zip(entries: dict[str, bytes]) -> bytes:
    """Build an in-memory ZIP with the given filename→bytes entries."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _mock_response(data: bytes, status: int = 200):
    resp = MagicMock()
    resp.content = data
    resp.raise_for_status = MagicMock()
    if status >= 400:
        resp.raise_for_status.side_effect = Exception(f"HTTP {status}")
    return resp


class ZipExtractorTest(unittest.TestCase):
    def _call(self, zip_data: bytes, verify_ssl: bool = True):
        with patch("app.extractors.zip_extractor.requests.get", return_value=_mock_response(zip_data)):
            return extract_text_from_zip("https://mshsk.ru/download/test.zip", verify_ssl=verify_ssl)

    def test_extracts_text_from_docx_inside_zip(self) -> None:
        from docx import Document as DocxDoc
        buf = io.BytesIO()
        doc = DocxDoc()
        doc.add_paragraph("Субсидия на развитие АПК")
        doc.save(buf)
        zip_data = _make_zip({"prikaz-subsid.docx": buf.getvalue()})

        result = self._call(zip_data)

        self.assertEqual(result.document_type, "zip")
        self.assertIn("Субсидия на развитие АПК", result.raw_text)

    def test_nested_zip_is_skipped(self) -> None:
        inner_zip = _make_zip({"inner.txt": b"should not appear"})
        outer_zip = _make_zip({"nested.zip": inner_zip})

        result = self._call(outer_zip)

        self.assertEqual(result.raw_text, "")
        self.assertIsNotNone(result.error)

    def test_macos_metadata_entries_are_skipped(self) -> None:
        from docx import Document as DocxDoc
        buf = io.BytesIO()
        doc = DocxDoc()
        doc.add_paragraph("Реальный документ")
        doc.save(buf)
        zip_data = _make_zip({
            "__MACOSX/._prikaz.docx": b"macos garbage",
            "prikaz.docx": buf.getvalue(),
        })

        result = self._call(zip_data)

        self.assertIn("Реальный документ", result.raw_text)

    def test_rejects_oversized_download(self) -> None:
        oversized = b"x" * (MAX_ZIP_DOWNLOAD_BYTES + 1)
        with patch("app.extractors.zip_extractor.requests.get", return_value=_mock_response(oversized)):
            result = extract_text_from_zip("https://mshsk.ru/download/big.zip")

        self.assertEqual(result.raw_text, "")
        self.assertIsNotNone(result.error)
        self.assertIn("too large", result.error)

    def test_rejects_bad_zip(self) -> None:
        result = self._call(b"not a zip")

        self.assertEqual(result.raw_text, "")
        self.assertIsNotNone(result.error)
        self.assertIn("bad format", result.error.lower())

    def test_zip_with_no_extractable_entries_returns_error(self) -> None:
        zip_data = _make_zip({"readme.txt": b"plain text file"})

        result = self._call(zip_data)

        self.assertEqual(result.raw_text, "")
        self.assertIsNotNone(result.error)
        self.assertIn("no extractable text", result.error)

    def test_fetch_error_returns_error_result(self) -> None:
        with patch("app.extractors.zip_extractor.requests.get", side_effect=Exception("connection refused")):
            result = extract_text_from_zip("https://mshsk.ru/download/fail.zip")

        self.assertEqual(result.raw_text, "")
        self.assertIsNotNone(result.error)
        self.assertIn("fetch error", result.error.lower())

    def test_multiple_docx_entries_are_concatenated(self) -> None:
        from docx import Document as DocxDoc

        def _docx_bytes(text: str) -> bytes:
            buf = io.BytesIO()
            d = DocxDoc()
            d.add_paragraph(text)
            d.save(buf)
            return buf.getvalue()

        zip_data = _make_zip({
            "doc1.docx": _docx_bytes("Грантовая поддержка"),
            "doc2.docx": _docx_bytes("Субсидия на мелиорацию"),
        })

        result = self._call(zip_data)

        self.assertIn("Грантовая поддержка", result.raw_text)
        self.assertIn("Субсидия на мелиорацию", result.raw_text)

    def test_verify_ssl_false_passes_verify_false_to_requests(self) -> None:
        zip_data = _make_zip({"readme.txt": b"irrelevant"})
        with patch("app.extractors.zip_extractor.requests.get", return_value=_mock_response(zip_data)) as mock_get:
            extract_text_from_zip("https://mshsk.ru/download/test.zip", verify_ssl=False)
        mock_get.assert_called_once()
        _, call_kwargs = mock_get.call_args
        self.assertFalse(call_kwargs.get("verify", True))

    def test_xls_inside_zip_produces_metadata_fallback(self) -> None:
        zip_data = _make_zip({"data.xls": b"\xd0\xcf\x11\xe0fake-xls-binary"})

        result = self._call(zip_data)

        self.assertEqual(result.document_type, "zip")
        self.assertIn("data.xls", result.raw_text)
        self.assertIn("xlrd", result.raw_text)


if __name__ == "__main__":
    unittest.main()
