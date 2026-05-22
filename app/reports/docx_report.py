from __future__ import annotations

import re
from pathlib import Path

from docx import Document


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*[-*]\s+(.+?)\s*$")
_NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+(.+?)\s*$")
_MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)]\(([^)]+)\)")
_BOLD_ITALIC_MARK_RE = re.compile(r"(\*\*|__|\*|_)")


def create_docx_from_markdown_file(
    markdown_path: Path | str,
    output_path: Path | str | None = None,
) -> Path:
    source_path = Path(markdown_path)
    target_path = Path(output_path) if output_path is not None else source_path.with_suffix(".docx")
    markdown = source_path.read_text(encoding="utf-8-sig")
    return create_docx_from_markdown(markdown, target_path)


def create_docx_from_markdown(markdown: str, output_path: Path | str) -> Path:
    target_path = Path(output_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document)

    previous_blank = False
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            if not previous_blank:
                document.add_paragraph("")
            previous_blank = True
            continue
        previous_blank = False

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            document.add_heading(_clean_inline_markdown(heading_match.group(2)), level=level)
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            document.add_paragraph(_clean_inline_markdown(bullet_match.group(1)), style="List Bullet")
            continue

        numbered_match = _NUMBERED_RE.match(line)
        if numbered_match:
            document.add_paragraph(_clean_inline_markdown(numbered_match.group(1)), style="List Number")
            continue

        document.add_paragraph(_clean_inline_markdown(line))

    document.save(target_path)
    return target_path


def _configure_document(document: Document) -> None:
    try:
        normal = document.styles["Normal"]
    except KeyError:
        return
    normal.font.name = "Arial"


def _clean_inline_markdown(text: str) -> str:
    normalized = _MARKDOWN_LINK_RE.sub(_visible_markdown_link, text)
    normalized = _BOLD_ITALIC_MARK_RE.sub("", normalized)
    normalized = normalized.replace("`", "")
    return normalized.strip()


def _visible_markdown_link(match: re.Match[str]) -> str:
    label = match.group(1).strip()
    url = match.group(2).strip()
    if not label:
        return url
    if label == url:
        return url
    return f"{label} ({url})"
